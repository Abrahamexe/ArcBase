"""Metni Word olarak indirme: baslik, bilgi tablosu, not satiri, OCR
paragraflari, kontrol karakterleri, veri ayrimi."""

import io

from docx import Document

import database
import word_uret
from conftest import belge_ekle


def _docx(c, belge_id):
    r = c.get(f"/belge/{belge_id}/word")
    assert r.status_code == 200 and r.mimetype == word_uret.MIME
    return r, Document(io.BytesIO(r.data))


def test_word_icerigi(girisli):
    c, k = girisli
    metin = "E-ARSIV FATURA\nABC Elektronik\n\nFatura No: ABC123\nGenel Toplam: 1.200,00 TL\x0c"
    belge_id = belge_ekle(k["id"], "fatura.jpg", "fatura", metin, kurum="ABC Elektronik",
                          tutar="1.200,00", para_birimi="TRY", belge_tarihi="2026-03-15")
    r, doc = _docx(c, belge_id)
    assert r.headers["Content-Disposition"] == "attachment; filename=fatura_2026-03-15.docx"

    paragraflar = [p.text for p in doc.paragraphs]
    assert paragraflar[0] == "Fatura"  # baslik: turun gorunen adi
    assert doc.paragraphs[0].style.name == "Title"
    assert "fatura.jpg · 15.03.2026" in paragraflar
    assert word_uret.NOT_METNI in paragraflar

    tablo = {satir.cells[0].text: satir.cells[1].text for satir in doc.tables[0].rows}
    assert tablo["Firma"] == "ABC Elektronik" and tablo["Tutar"] == "1.200,00 TL"
    assert tablo["Fatura tarihi"] == "15.03.2026" and tablo["Fatura no"] == "—"

    # OCR metni: bos satirla ayrilan bloklar ayri paragraf, satir sonlari korunur
    assert "E-ARSIV FATURA\nABC Elektronik" in paragraflar
    assert "Fatura No: ABC123\nGenel Toplam: 1.200,00 TL" in paragraflar  # \x0c temizlendi
    assert doc.core_properties.author == "ArcBase"


def test_metinsiz_ve_eski_belge(girisli):
    c, k = girisli
    belge_id = belge_ekle(k["id"], "bos.jpg", "diger", "")  # v3: bilgi alanlari NULL
    r, doc = _docx(c, belge_id)
    assert "Bu belgeden metin okunamadı." in [p.text for p in doc.paragraphs]
    assert [s.cells[0].text for s in doc.tables[0].rows] == ["Konu", "Tarih"]
    assert r.headers["Content-Disposition"].startswith("attachment; filename=diger_20")


def test_word_veri_ayrimi(girisli, kullanici, istemci):
    c, _ = girisli
    diger = kullanici()
    yabanci = belge_ekle(diger["id"], "gizli.jpg", "fatura", "GIZLI METIN")
    assert c.get(f"/belge/{yabanci}/word").status_code == 404
    assert istemci.get(f"/belge/{yabanci}/word").status_code == 302
    assert database.belge_getir(diger["id"], yabanci) is not None


def test_detay_sayfasinda_word_butonu(girisli):
    c, k = girisli
    belge_id = belge_ekle(k["id"])
    assert f'href="/belge/{belge_id}/word"' in c.get(f"/belge/{belge_id}").get_data(as_text=True)
