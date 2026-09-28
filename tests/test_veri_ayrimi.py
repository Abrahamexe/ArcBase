"""Iki kullanici arasinda veri ayrimi: biri digerinin belgelerine hicbir
yoldan erisemez; erisim denemeleri 404 doner (belgenin varligi belli edilmez)."""

import pytest

import database
from conftest import belge_ekle, giris_yap


@pytest.fixture
def iki_kullanici(uygulama, kullanici):
    a, b = kullanici(), kullanici()
    a["belge"] = belge_ekle(a["id"], "a_gizli_fatura.jpg", "fatura", "ALFAGIZLI metin")
    b["belge"] = belge_ekle(b["id"], "b_sozlesme.jpg", "sozlesme", "BETAMETIN")
    a["istemci"], b["istemci"] = uygulama.test_client(), uygulama.test_client()
    giris_yap(a["istemci"], a["ad"])
    giris_yap(b["istemci"], b["ad"])
    return a, b


@pytest.mark.parametrize(
    "yol", ["/belge/{id}", "/gorsel/{id}", "/onizleme/{id}", "/belge/{id}/indir"]
)
def test_baskasinin_belgesine_get_404(iki_kullanici, yol):
    a, b = iki_kullanici
    r = b["istemci"].get(yol.format(id=a["belge"]))
    assert r.status_code == 404
    assert "a_gizli_fatura" not in r.get_data(as_text=True)
    # Sahibi erisebilir
    assert a["istemci"].get(yol.format(id=a["belge"])).status_code == 200


def test_baskasinin_belgesi_ile_olmayan_belge_ayni_yanit(iki_kullanici):
    a, b = iki_kullanici
    yabanci = b["istemci"].get(f"/belge/{a['belge']}")
    olmayan = b["istemci"].get("/belge/2147480000")
    assert yabanci.status_code == olmayan.status_code == 404
    assert yabanci.get_data() == olmayan.get_data()


def test_baskasinin_belgesini_degistiremez_silemez(iki_kullanici):
    a, b = iki_kullanici
    assert b["istemci"].post(f"/belge/{a['belge']}/tur", data={"tur": "diger"}).status_code == 404
    assert b["istemci"].post(f"/belge/{a['belge']}/sil").status_code == 404
    kayit = database.belge_getir(a["id"], a["belge"])
    assert kayit is not None and kayit["tur"] == "fatura" and not kayit["elle_duzeltildi"]


def test_listeler_arama_ve_sayaclar_ayri(iki_kullanici):
    a, b = iki_kullanici
    c = b["istemci"]
    panel = c.get("/").get_data(as_text=True)
    assert "b_sozlesme.jpg" in panel and "a_gizli_fatura" not in panel
    assert "a_gizli_fatura" not in c.get("/kategori/fatura").get_data(as_text=True)
    assert "a_gizli_fatura" not in c.get("/ara?q=ALFAGIZLI").get_data(as_text=True)
    assert "a_gizli_fatura" not in c.get("/ara?q=a_gizli").get_data(as_text=True)
    assert "b_sozlesme" in c.get("/ara?q=BETAMETIN").get_data(as_text=True)

    sayilar = database.tur_sayilari(b["id"])
    assert sayilar["fatura"]["adet"] == 0 and sayilar["sozlesme"]["adet"] == 1
    assert database.belge_sayisi(b["id"]) == 1
    assert "/ 50 belge" in c.get("/hesap").get_data(as_text=True)


def test_veritabani_katmani_kullanici_filtresi(iki_kullanici):
    """database.py'deki her belge fonksiyonu baska kullanicinin id'siyle bos doner."""
    a, b = iki_kullanici
    assert database.belge_getir(b["id"], a["belge"]) is None
    assert database.belge_icerigi(b["id"], a["belge"]) is None
    assert database.belge_icerigi(b["id"], a["belge"], onizleme=True) is None
    assert all(x["id"] != a["belge"] for x in database.belgeleri_listele(b["id"]))
    assert database.belgeleri_suz(b["id"], tur="fatura") == []
    assert database.belgeleri_suz(b["id"], sorgu="ALFAGIZLI") == []
    assert database.turu_guncelle(b["id"], a["belge"], "diger") is False
    assert database.belge_sil(b["id"], a["belge"]) is False
    assert database.belge_getir(a["id"], a["belge"])["tur"] == "fatura"


def test_hesap_silinince_belgeler_de_silinir(iki_kullanici):
    a, b = iki_kullanici
    database.kullanici_sil(a["id"])
    assert database.belge_getir(a["id"], a["belge"]) is None
    assert database.belge_getir(b["id"], b["belge"]) is not None


def test_v4_uc_noktalarinda_veri_ayrimi(iki_kullanici):
    """Bilgi duzenleme, filtre/arama (kurum, alanlar), Excel ve toplu bilgi
    cikarma baska kullanicinin belgelerine ulasmaz."""
    import io

    from openpyxl import load_workbook

    import alanlar

    a, b = iki_kullanici
    database.alanlari_guncelle(a["id"], a["belge"], alanlar.dogrula(
        "fatura", {"kurum": "ALFAFIRMA", "tutar": 50, "fatura_no": "ALFA-NO-1"})[0])
    c = b["istemci"]

    assert c.post(f"/belge/{a['belge']}/alanlar", data={"kurum": "ele gecirildi"}).status_code == 404
    assert database.belge_getir(a["id"], a["belge"])["kurum"] == "ALFAFIRMA"

    for adres in ("/ara?q=ALFAFIRMA", "/ara?q=ALFA-NO-1", "/kategori/fatura?kurum=ALFA",
                  "/kategori/fatura?min_tutar=1"):
        assert "a_gizli_fatura" not in c.get(adres).get_data(as_text=True), adres

    for adres in ("/disa-aktar", "/disa-aktar?tur=fatura", "/disa-aktar?q=ALFAFIRMA"):
        sayfa = load_workbook(io.BytesIO(c.get(adres).data)).active
        degerler = [h.value for satir in sayfa.iter_rows() for h in satir]
        assert "a_gizli_fatura.jpg" not in degerler and "ALFAFIRMA" not in degerler, adres

    # b'nin toplu bilgi cikarmasi a'nin v3'ten kalan belgesine dokunmaz
    database._sorgula("UPDATE belgeler SET alanlar = NULL, kurum = NULL WHERE id = %s", (a["belge"],))
    r = c.post("/hesap/bilgi-cikar", headers={"Accept": "application/json"}).get_json()
    assert r["son_id"] != a["belge"]
    assert database.belge_getir(a["id"], a["belge"])["alanlar"] is None
    assert database.bilgisi_eksik_sayisi(b["id"]) == 0
