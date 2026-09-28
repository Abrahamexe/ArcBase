"""Excel'e aktarma: dosya adi, basliklar, hucre tipleri, suzgecler, toplam
satiri, veri ayrimi ve formul enjeksiyonuna karsi koruma."""

import io
import re
from datetime import date, datetime
from decimal import Decimal

import pytest
from openpyxl import load_workbook

import database
import excel
from conftest import belge_ekle

BASLIKLAR = [
    "Dosya adı", "Tür", "Belge tarihi", "Kurum", "Tutar", "Para birimi",
    "Güven", "Elle düzeltildi", "Yüklenme tarihi", "Açıklama",
]


def indir(c, sorgu=""):
    r = c.get(f"/disa-aktar?{sorgu}")
    assert r.status_code == 200 and r.mimetype == excel.MIME
    return r, load_workbook(io.BytesIO(r.data)).active


def satirlar(sayfa):
    return [[h.value for h in satir] for satir in sayfa.iter_rows(min_row=2)]


@pytest.fixture
def arsiv(girisli, kullanici):
    c, k = girisli
    belge_ekle(k["id"], "ocak.jpg", "fatura", "FATURA", 0.95,
               kurum="Enerjisa", tutar="1.234,50", para_birimi="TRY", belge_tarihi="2026-01-10")
    belge_ekle(k["id"], "subat.jpg", "fatura", "FATURA", 0.9,
               kurum="İSKİ", tutar=300, para_birimi="TRY", belge_tarihi="2026-02-10")
    belge_ekle(k["id"], "aws.jpg", "fatura", "FATURA", 0.9,
               kurum="AWS", tutar=20, para_birimi="USD", belge_tarihi="2026-02-20")
    belge_ekle(k["id"], "dilekce.jpg", "dilekce", "belediye", 0.8, kurum="Kadıköy Belediyesi", konu="Park")
    baskasi = kullanici()
    belge_ekle(baskasi["id"], "gizli.jpg", "fatura", "FATURA", kurum="Gizli A.Ş.", tutar=999)
    return c, k


def test_fatura_disa_aktarma(arsiv):
    c, k = arsiv
    r, sayfa = indir(c, "tur=fatura")
    assert re.fullmatch(r'attachment; filename=arsiv_fatura_\d{4}-\d{2}-\d{2}\.xlsx', r.headers["Content-Disposition"])

    assert [h.value for h in sayfa[1]] == BASLIKLAR
    assert all(h.font.bold for h in sayfa[1]) and sayfa.freeze_panes == "A2"

    veri = satirlar(sayfa)
    belgeler = [s for s in veri if s[0] and s[0].endswith(".jpg")]
    assert [s[0] for s in belgeler] == ["aws.jpg", "subat.jpg", "ocak.jpg"]  # gizli.jpg yok
    ocak = belgeler[2]
    assert ocak[1] == "Fatura" and ocak[3] == "Enerjisa" and ocak[5] == "TRY"
    assert ocak[2] == datetime(2026, 1, 10)  # metin degil, gercek tarih
    assert isinstance(ocak[4], (int, float)) and ocak[4] == pytest.approx(1234.5)
    assert ocak[6] == pytest.approx(0.95) and ocak[7] == "Hayır"
    assert isinstance(ocak[8], datetime) and ocak[9] == "test"

    ocak_satiri = sayfa[4]
    assert ocak_satiri[2].number_format == excel.TARIH_BICIMI
    assert ocak_satiri[4].number_format == excel.TUTAR_BICIMI
    assert ocak_satiri[6].number_format == excel.GUVEN_BICIMI

    # En altta para birimi basina toplam
    toplamlar = {s[5]: s[4] for s in veri if s[0] == "Toplam"}
    assert toplamlar == {"TRY": pytest.approx(1534.5), "USD": pytest.approx(20)}
    son = sayfa[sayfa.max_row]
    assert son[0].value == "Toplam" and son[0].font.bold

    # sutun genislikleri icerige gore
    assert sayfa.column_dimensions["A"].width >= len("Dosya adı")
    assert sayfa.column_dimensions["J"].width < excel.MAKS_SUTUN_GENISLIGI + 1


def test_suzgecler_aynen_uygulanir(arsiv):
    c, _ = arsiv
    _, sayfa = indir(c, "tur=fatura&baslangic=2026-02-01&bitis=2026-02-15")
    veri = satirlar(sayfa)
    assert [s[0] for s in veri if s[0] and s[0].endswith(".jpg")] == ["subat.jpg"]
    assert [s[4] for s in veri if s[0] == "Toplam"] == [300]

    _, sayfa = indir(c, "tur=fatura&siralama=tutar&min_tutar=100")
    assert [s[0] for s in satirlar(sayfa) if s[0] and s[0].endswith(".jpg")] == ["ocak.jpg", "subat.jpg"]

    _, sayfa = indir(c, "q=kadıköy")
    assert [s[0] for s in satirlar(sayfa)] == ["dilekce.jpg"]


def test_tum_arsiv_ve_toplamsiz(arsiv):
    c, _ = arsiv
    r, sayfa = indir(c)
    assert "arsiv_tumu_" in r.headers["Content-Disposition"]
    veri = satirlar(sayfa)
    assert len(veri) == 4 and "Toplam" not in [s[0] for s in veri]
    assert "gizli.jpg" not in [s[0] for s in veri]


def test_elle_duzeltme_ve_bos_alanlar(girisli):
    c, k = girisli
    belge_id = belge_ekle(k["id"], "eski.jpg", "kimlik")  # v3: bilgi alanlari NULL
    database.turu_guncelle(k["id"], belge_id, "resmi_yazi")
    _, sayfa = indir(c)
    satir = satirlar(sayfa)[0]
    assert satir[1] == "Resmi Yazı" and satir[7] == "Evet"
    assert satir[2] is None and satir[4] is None


def test_formul_enjeksiyonu_calistirilmaz():
    belge = {
        "dosya_adi": "rapor.jpg", "tur": "fatura", "belge_tarihi": date(2026, 1, 1),
        "kurum": '=HYPERLINK("http://ornek.com","tikla")', "tutar": Decimal("5"), "para_birimi": "TRY",
        "guven": 0.5, "elle_duzeltildi": False, "alanlar_elle_duzeltildi": False,
        "yuklenme_tarihi": datetime(2026, 1, 1, 9, 0), "aciklama": "kontrol\x0bkarakteri",
    }
    sayfa = load_workbook(io.BytesIO(excel.olustur([belge], database.TURLER))).active
    kurum = sayfa["D2"]
    assert kurum.data_type == "s" and kurum.value.startswith("=HYPERLINK")
    assert sayfa["J2"].value == "kontrolkarakteri"
    assert sayfa["I2"].value == datetime(2026, 1, 1, 12, 0)  # UTC -> TSI


def test_baglantilar_sayfalarda(arsiv):
    c, _ = arsiv
    kategori = c.get("/kategori/fatura?kurum=Enerjisa").get_data(as_text=True)
    assert "/disa-aktar?kurum=Enerjisa&amp;tur=fatura" in kategori
    arama = c.get("/ara?q=belediye&tur=dilekce").get_data(as_text=True)
    assert "/disa-aktar?q=belediye&amp;tur=dilekce" in arama
    assert 'href="/disa-aktar"' in c.get("/").get_data(as_text=True)


def test_giris_gerekli(istemci):
    r = istemci.get("/disa-aktar")
    assert r.status_code == 302 and "/giris" in r.headers["Location"]
