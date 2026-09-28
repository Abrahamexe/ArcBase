"""Belge listesini ve cikarilan bilgileri Excel (.xlsx) dosyasina donusturme.

OCR metni ve gorseller disa aktarilmaz. Tarih ve tutar hucreleri metin degil,
gercek tarih/sayi olarak yazilir; Excel'de siralanip toplanabilir.
"""

import io
from datetime import timedelta

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill

MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
# Veritabani saati UTC; Turkiye 2016'dan beri yaz saati uygulamadan sabit UTC+3.
TSI_FARKI = timedelta(hours=3)

TARIH_BICIMI = "DD.MM.YYYY"
TARIH_SAAT_BICIMI = "DD.MM.YYYY HH:MM"
TUTAR_BICIMI = "#,##0.00"
GUVEN_BICIMI = "0%"
MAKS_SUTUN_GENISLIGI = 60


def _sutunlar(turler):
    """(baslik, deger_fonksiyonu, hucre_bicimi)"""
    return (
        ("Dosya adı", lambda b: b["dosya_adi"], None),
        ("Tür", lambda b: turler.get(b["tur"], {}).get("ad", b["tur"]), None),
        ("Belge tarihi", lambda b: b["belge_tarihi"], TARIH_BICIMI),
        ("Kurum", lambda b: b["kurum"], None),
        ("Tutar", lambda b: b["tutar"], TUTAR_BICIMI),
        ("Para birimi", lambda b: b["para_birimi"], None),
        ("Güven", lambda b: b["guven"], GUVEN_BICIMI),
        (
            "Elle düzeltildi",
            lambda b: "Evet" if b["elle_duzeltildi"] or b["alanlar_elle_duzeltildi"] else "Hayır",
            None,
        ),
        ("Yüklenme tarihi", lambda b: b["yuklenme_tarihi"] + TSI_FARKI, TARIH_SAAT_BICIMI),
        ("Açıklama", lambda b: b["aciklama"], None),
    )


def _hucre_degeri(deger):
    if isinstance(deger, str):
        # OCR/yapay zeka metninde Excel'in kabul etmedigi kontrol karakterleri olabilir.
        return ILLEGAL_CHARACTERS_RE.sub("", deger)
    return deger


def _gorunen_uzunluk(deger, bicim):
    if deger is None:
        return 0
    if bicim == TARIH_BICIMI:
        return 10
    if bicim == TARIH_SAAT_BICIMI:
        return 16
    if bicim == TUTAR_BICIMI:
        return len(f"{deger:,.2f}")
    if bicim == GUVEN_BICIMI:
        return 4
    return max(len(satir) for satir in str(deger).splitlines() or [""])


def olustur(belgeler, turler, tutar_toplamlari=None):
    """belgeler: database.belgeleri_suz satirlari. tutar_toplamlari verilirse
    ({'TRY': Decimal, ...}) en alta para birimi basina toplam satiri eklenir."""
    sutunlar = _sutunlar(turler)
    kitap = Workbook()
    sayfa = kitap.active
    sayfa.title = "Belgeler"

    sayfa.append([baslik for baslik, _, _ in sutunlar])
    for hucre in sayfa[1]:
        hucre.font = Font(bold=True, color="FFFFFF")
        hucre.fill = PatternFill("solid", fgColor="0061FE")
        hucre.alignment = Alignment(vertical="center")
    sayfa.freeze_panes = "A2"

    genislikler = [len(baslik) for baslik, _, _ in sutunlar]
    for belge in belgeler:
        sayfa.append([_hucre_degeri(deger(belge)) for _, deger, _ in sutunlar])
        for i, (hucre, (_, _, bicim)) in enumerate(zip(sayfa[sayfa.max_row], sutunlar)):
            if isinstance(hucre.value, str) and hucre.value.startswith("="):
                hucre.data_type = "s"  # dosya adi/kurum formul olarak calistirilmasin
            if bicim:
                hucre.number_format = bicim
            genislikler[i] = max(genislikler[i], _gorunen_uzunluk(hucre.value, bicim))

    if tutar_toplamlari:
        sayfa.append([])
        for para, toplam in sorted(tutar_toplamlari.items()):
            sayfa.append(["Toplam", None, None, None, toplam, para or None])
            for hucre in sayfa[sayfa.max_row]:
                hucre.font = Font(bold=True)
            sayfa.cell(row=sayfa.max_row, column=5).number_format = TUTAR_BICIMI

    for i, genislik in enumerate(genislikler, start=1):
        sayfa.column_dimensions[sayfa.cell(row=1, column=i).column_letter].width = min(
            genislik + 2, MAKS_SUTUN_GENISLIGI
        )

    tampon = io.BytesIO()
    kitap.save(tampon)
    return tampon.getvalue()
