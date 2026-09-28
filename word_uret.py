"""Belge metnini Word (.docx) olarak uretme (python-docx).

Icerik: baslik (turun gorunen adi), dosya adi ve tarih, cikarilan bilgi
alanlari kucuk bir tablo halinde, "otomatik okunmustur" notu ve OCR metni
paragraflar halinde. Saklanmaz, her istekte uretilir.
"""

import io
import re

from docx import Document
from docx.enum.text import WD_BREAK
from docx.shared import Pt, RGBColor

import alanlar

MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
NOT_METNI = "Bu metin otomatik okunmuştur, hatalar içerebilir."
URETICI = "ArcBase"
# XML'de gecersiz kontrol karakterleri (OCR ciktisinda form feed vb. olabilir)
_GECERSIZ = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f]")


def _temiz(metin):
    return _GECERSIZ.sub("", metin or "")


def _deger_yaz(tip, deger, para_birimi):
    if deger is None or deger == "":
        return "—"
    if tip == "tarih":
        return alanlar.tarih_yaz(deger)
    if tip == "tutar":
        return alanlar.tutar_yaz(deger, para_birimi)
    return _temiz(str(deger))


def olustur(belge, tur_adi, alt_baslik):
    """belge: database.belge_getir kaydi. Donus: .docx baytlari."""
    doc = Document()
    doc.core_properties.title = tur_adi
    doc.core_properties.author = URETICI
    doc.core_properties.comments = NOT_METNI

    doc.add_heading(tur_adi, level=0)
    alt = doc.add_paragraph(_temiz(alt_baslik))
    alt.runs[0].font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    degerler = alanlar.degerler(belge)
    tanimlar = alanlar.ALANLAR.get(belge["tur"], alanlar.ALANLAR["diger"])
    doc.add_heading("Belge bilgileri", level=2)
    tablo = doc.add_table(rows=0, cols=2)
    tablo.style = "Light Grid Accent 1"
    for anahtar, etiket, tip in tanimlar:
        hucreler = tablo.add_row().cells
        hucreler[0].text = etiket
        hucreler[0].paragraphs[0].runs[0].font.bold = True
        hucreler[1].text = _deger_yaz(tip, degerler.get(anahtar), degerler.get("para_birimi"))

    doc.add_heading("Metin", level=2)
    not_paragrafi = doc.add_paragraph()
    not_yazisi = not_paragrafi.add_run(NOT_METNI)
    not_yazisi.italic = True
    not_yazisi.font.size = Pt(9)
    not_yazisi.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    metin = _temiz(belge.get("ocr_metni")).strip()
    if not metin:
        doc.add_paragraph("Bu belgeden metin okunamadı.")
    # Bos satirla ayrilan bloklar ayri paragraf; blok icindeki satirlar satir sonuyla korunur.
    for blok in re.split(r"\n\s*\n", metin):
        satirlar = [s.rstrip() for s in blok.strip().splitlines()]
        if not satirlar:
            continue
        paragraf = doc.add_paragraph()
        for i, satir in enumerate(satirlar):
            yazi = paragraf.add_run(satir)
            if i < len(satirlar) - 1:
                yazi.add_break(WD_BREAK.LINE)

    tampon = io.BytesIO()
    doc.save(tampon)
    return tampon.getvalue()
